"""
Repair Order Matcher
====================
Matches the Repair Order numbers on Sheet 1 (an .xlsx file) against the
numbers on Sheet 2 and Sheet 3 (.xls or .xlsx), subtracts the
Sheet 2/3 Balance from the Sheet 1 Total Paid Amount, and writes the
result into a new column right of "Repair Order #" on Sheet 1.

    Positive result  -> pastel green
    Negative result  -> pastel red
    Zero             -> shown as *
    Not found        -> yellow, "Not found"

Before saving, a copy of Sheet 1 is made in the same
folder. Running it again on the same file updates the existing column
instead of adding another one.

How to run
    Double-click (or run with no arguments) to pick the files in pop-up windows.
    Or from a command prompt:
        python repair_order_matcher.py "Sheet 1.xlsx" "Sheet 2.xls" "Sheet 3.xls"

Needs:  pip install openpyxl xlrd

Packaging as a Windows .exe
    RepairOrderMatcher.exe is built with PyInstaller, which bundles this
    script, Python (3.12),openpyxl, xlrd and tkinter
    libraries into one file, so it runs on PCs that don't have Python
    installed. To rebuild it after changing this script:

        pip install pyinstaller openpyxl xlrd
        pyinstaller --onefile --windowed --name RepairOrderMatcher repair_order_matcher.py

    --onefile   packs everything into a single .exe. When it starts, it
                unpacks itself to a temporary folder, which is why the first
                launch takes a few seconds.
    --windowed  hides the black console window, so the user only sees the
                file pickers and message boxes. (This is also why xlrd's
                console warnings are redirected; see read_two_columns.)
    --name      sets the name of the .exe.

    The finished file appears in the dist\\ folder. The build\\ folder and
    the RepairOrderMatcher.spec file are build leftovers and can be deleted
    or left out.
"""

import io
import os
import re
import shutil
import sys
from copy import copy
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange

# ===========================================================================
# SETTINGS
# Column letters and colors used throughout the script. If the layout of
# a report ever changes, update the letters here; nothing else needs editing.
# ===========================================================================

# Sheet 1: the workbook that gets updated
SHEET1_RO_COLUMN = "B"          # "Repair Order #"
SHEET1_PRICE_COLUMN = "N"       # "Total Paid Amount"
SHEET1_HEADER_ROW = 1           # data starts on the row below this

# Sheet 2 / Sheet 3: the reports the repair orders are looked up in
LOOKUP_RO_COLUMN = "C"          # "Control"
LOOKUP_BALANCE_COLUMN = "V"     # "Balance"

# The result column added to Sheet 1
NEW_COLUMN_HEADER = "Price Difference"
NEW_COLUMN_WIDTH = 18

# Cell colors for the result column
GREEN = PatternFill(fill_type="solid", start_color="C6EFCE", end_color="C6EFCE")    # positive
RED = PatternFill(fill_type="solid", start_color="FFC7CE", end_color="FFC7CE")      # negative
YELLOW = PatternFill(fill_type="solid", start_color="FFFF99", end_color="FFFF99")   # needs attention
NO_FILL = PatternFill(fill_type=None)

# Currency with the full amount shown; negatives get a leading minus, zero displays as *
RESULT_NUMBER_FORMAT = '$#,##0.00;-$#,##0.00;"*"'


class MatcherError(Exception):
    """A problem we can explain to the user in plain words."""


# ===========================================================================
# VALUE CLEANUP
# Turns the raw contents of a cell into something that can be compared or
# calculated with, no matter how Excel or the report stored it.
# ===========================================================================
def normalize_ro(value):
    """Turn a Repair Order / Control number into a key that can be compared.

    Handles the usual reasons identical-looking numbers fail to match:
    number vs text (249569 vs "249569"), stray or missing spaces
    ("01293 7" vs "012937"), and leading zeros Excel dropped.
    """
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    key = re.sub(r"\s+", "", str(value)).upper()
    if key.isdigit():
        key = key.lstrip("0") or "0"
    return key


def to_number(value):
    """Read a price that may be a real number or text like "$1,250.00" or "(45.00)"."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace(",", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


def money(amount):
    """Format a number as currency for notes and messages, e.g. -$45.00."""
    return f"-${abs(amount):,.2f}" if amount < 0 else f"${amount:,.2f}"


# ===========================================================================
# READING SHEET 2 / SHEET 3
# Pulls the RO number and Balance from every row of a lookup report
# (old .xls or newer .xlsx) and builds a table to look repair orders up in.
# ===========================================================================
def read_two_columns(path, key_column, value_column):
    """Yield (key, value) from the given columns of every row of every worksheet.

    The reader is chosen from the file's contents, not its extension, so an
    .xlsx that was renamed to .xls (or the other way round) still works.
    """
    k = column_index_from_string(key_column) - 1
    v = column_index_from_string(value_column) - 1
    with open(path, "rb") as f:
        signature = f.read(8)

    if signature.startswith(b"PK"):  # .xlsx
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    yield (row[k] if k < len(row) else None,
                           row[v] if v < len(row) else None)
        finally:
            wb.close()

    elif signature.startswith(b"\xD0\xCF\x11\xE0"):  # old .xls
        try:
            import xlrd
        except ImportError:
            raise MatcherError("Reading .xls files needs the xlrd package.\n"
                               "Install it with:  pip install xlrd")
        # xlrd prints warnings to the console by default; a packaged .exe has
        # no console, so send them somewhere harmless.
        log = io.StringIO()
        try:
            # Crystal Reports exports often trip xlrd's corruption check even
            # though the data is fine.
            book = xlrd.open_workbook(path, logfile=log, ignore_workbook_corruption=True)
        except TypeError:  # older xlrd without that option
            book = xlrd.open_workbook(path, logfile=log)
        usable = (xlrd.XL_CELL_TEXT, xlrd.XL_CELL_NUMBER)
        for sheet in book.sheets():
            for r in range(sheet.nrows):
                row = sheet.row(r)
                key = row[k] if k < len(row) else None
                val = row[v] if v < len(row) else None
                yield (key.value if key is not None and key.ctype in usable else None,
                       val.value if val is not None and val.ctype in usable else None)
    else:
        raise MatcherError(f"{os.path.basename(path)} doesn't look like an Excel file (.xls or .xlsx).")


def build_lookup(path):
    """Return {repair order key: [balance, balance, ...]} for one Sheet 2/3 file."""
    lookup = {}
    for ro_value, balance_value in read_two_columns(path, LOOKUP_RO_COLUMN, LOOKUP_BALANCE_COLUMN):
        key = normalize_ro(ro_value)
        balance = to_number(balance_value)
        if key and balance is not None:
            lookup.setdefault(key, []).append(balance)
    if not lookup:
        raise MatcherError(
            f"No control numbers with balances were found in columns "
            f"{LOOKUP_RO_COLUMN} and {LOOKUP_BALANCE_COLUMN} of:\n{path}\n\n"
            f"Check that this is the right file. No changes were made.")
    return lookup


# ===========================================================================
# checks the files, matches each repair order, fills in the
# result column, backs up the original and saves.
# ===========================================================================
def insert_column(ws, col_idx):
    """Insert an empty column at col_idx.

    openpyxl moves the cells but not the column widths or the filter range,
    so those are shifted here to stay lined up with their data.
    """
    ws.insert_cols(col_idx)

    # Shift column widths one place to the right
    moved = {}
    for key, dim in list(ws.column_dimensions.items()):
        lo = dim.min or column_index_from_string(key)
        hi = dim.max or lo
        if lo >= col_idx:
            lo, hi = lo + 1, hi + 1
        elif hi >= col_idx:
            hi += 1
        dim.min, dim.max = lo, hi
        dim.index = get_column_letter(lo)
        moved[dim.index] = dim
    ws.column_dimensions.clear()
    ws.column_dimensions.update(moved)

    # Shift or widen the filter range
    if ws.auto_filter.ref:
        rng = CellRange(ws.auto_filter.ref)
        if rng.min_col >= col_idx:
            rng.shift(col_shift=1)
        elif rng.max_col >= col_idx:
            rng.expand(right=1)
        ws.auto_filter.ref = rng.coord


def process(sheet1_path, lookup_files, confirm=lambda message: True):
    """Do the matching and update Sheet 1.

    lookup_files: list of (label, path), e.g. [("Sheet 2", "a.xls"), ("Sheet 3", "b.xls")].
                  If a repair order is on more than one, the first one listed is used.
    confirm:      called with a question when something looks off; return True to continue.
    Returns a summary to show the user.
    """
    # --- Step 1: Safety checks before anything is changed -----------------
    # Sheet 1 must be an .xlsx (the only format openpyxl can write) and must
    # not be locked by Excel.
    with open(sheet1_path, "rb") as f:
        if not f.read(2) == b"PK":
            raise MatcherError("Sheet 1 must be an .xlsx file so it can be edited.\n"
                               "Open it in Excel, use File > Save As > Excel Workbook (*.xlsx), "
                               "and pick that file instead.")
    try:
        open(sheet1_path, "r+b").close()
    except PermissionError:
        raise MatcherError(f"Can't write to:\n{sheet1_path}\n\n"
                           "If it's open in Excel, close it and try again.")

    # --- Step 2: Load the lookup reports and Sheet 1 ----------------------
    lookups = [(label, path, build_lookup(path)) for label, path in lookup_files]

    keep_vba = sheet1_path.lower().endswith(".xlsm")
    wb = load_workbook(sheet1_path, keep_vba=keep_vba)                    # for writing (keeps formulas)
    wb_values = load_workbook(sheet1_path, data_only=True, keep_vba=keep_vba)  # for reading values
    ws, ws_values = wb.worksheets[0], wb_values.worksheets[0]

    # Work out column positions before and after the new column is inserted
    hdr = SHEET1_HEADER_ROW
    ro_idx = column_index_from_string(SHEET1_RO_COLUMN)
    new_idx = ro_idx + 1
    price_idx = column_index_from_string(SHEET1_PRICE_COLUMN)
    price_idx_after = price_idx + 1 if price_idx >= new_idx else price_idx

    # If this file was already processed, reuse that column instead of adding another.
    already_done = str(ws.cell(hdr, new_idx).value or "").strip() == NEW_COLUMN_HEADER
    price_read_idx = price_idx_after if already_done else price_idx

    # --- Step 3: Confirm this looks like the right Sheet 1 ----------------
    ro_header = str(ws_values.cell(hdr, ro_idx).value or "")
    price_header = str(ws_values.cell(hdr, price_read_idx).value or "")
    problems = []
    if "repair order" not in ro_header.lower():
        problems.append(f'Column {SHEET1_RO_COLUMN} is headed "{ro_header}", not "Repair Order #".')
    if "total paid" not in price_header.lower():
        problems.append(f'The price column is headed "{price_header}", not "Total Paid Amount".')
    if problems and not confirm("\n".join(problems) + "\n\nIs this the right Sheet 1 file? Continue anyway?"):
        return "Cancelled. No changes were made."

    # --- Step 4: Read every repair order and price from Sheet 1 -----------
    # Done before inserting the column so the original positions are used.
    rows = []
    for r in range(hdr + 1, ws_values.max_row + 1):
        ro_raw = ws_values.cell(r, ro_idx).value
        rows.append((r, ro_raw, normalize_ro(ro_raw), ws_values.cell(r, price_read_idx).value))
    if not any(key for _, _, key, _ in rows):
        raise MatcherError(f"No repair order numbers were found in column {SHEET1_RO_COLUMN} of Sheet 1. "
                           "No changes were made.")

    # --- Step 5: Re use the "Price Difference" column -------------
    if not already_done:
        insert_column(ws, new_idx)
    new_letter = get_column_letter(new_idx)
    price_letter = get_column_letter(price_idx_after)

    header_cell = ws.cell(hdr, new_idx)
    header_cell._style = copy(ws.cell(hdr, ro_idx)._style)
    header_cell.value = NEW_COLUMN_HEADER
    ws.column_dimensions[new_letter].width = NEW_COLUMN_WIDTH

    # --- Step 6: Match each repair order and fill in the result -----------
    counts = {"checked": 0, "not_found": 0, "no_price": 0,
              "positive": 0, "negative": 0, "zero": 0, "notes": 0}
    matched = {label: 0 for label, _, _ in lookups}

    for r, ro_raw, key, price_raw in rows:
        cell = ws.cell(r, new_idx)
        cell.value, cell.fill, cell.comment = None, NO_FILL, None   # clear any earlier run
        if not key:
            continue
        counts["checked"] += 1
        cell._style = copy(ws.cell(r, ro_idx)._style)
        cell.number_format = RESULT_NUMBER_FORMAT
        cell.fill = NO_FILL

        # Not on any lookup report: mark yellow 
        found = [(label, path, lookup[key]) for label, path, lookup in lookups if key in lookup]
        if not found:
            cell.value, cell.fill = "Not found", YELLOW
            counts["not_found"] += 1
            continue

        # Use the first report it was found on; note duplicates and other matches
        label, path, balances = found[0]
        matched[label] += 1
        other_total = round(sum(balances), 2)
        notes = []
        if len(balances) > 1:
            notes.append(f"Appears {len(balances)} times on {label} ({os.path.basename(path)}). "
                         f"Balances added together: "
                         f"{' + '.join(money(b) for b in balances)} = {money(other_total)}")
        for other_label, other_path, other_balances in found[1:]:
            notes.append(f"Also on {other_label} ({os.path.basename(other_path)}) with balance "
                         f"{money(round(sum(other_balances), 2))}. {label} was used.")

        # Calculate Total Paid Amount - Balance and color the cell
        price = to_number(price_raw)
        if price is None:
            cell.value, cell.fill = "No Sheet 1 price", YELLOW
            counts["no_price"] += 1
        else:
            diff = round(price - other_total, 2)
            if isinstance(price_raw, (int, float)) and not isinstance(price_raw, bool):
                # A formula, so clicking the cell shows the Sheet 2/3 amount used.
                sign = "-" if other_total >= 0 else "+"
                cell.value = f"=ROUND({price_letter}{r}{sign}{abs(other_total):.2f},2)"
            else:
                cell.value = diff
            if diff > 0:
                cell.fill = GREEN
                counts["positive"] += 1
            elif diff < 0:
                cell.fill = RED
                counts["negative"] += 1
            else:
                counts["zero"] += 1   # number format shows it as *

        # Attach any notes as an Excel comment (red triangle in the corner)
        if notes:
            cell.comment = Comment("\n".join(notes), "Repair Order Matcher", width=300, height=110)
            counts["notes"] += 1

    # --- Step 7: Back up the original, then save --------------------------
    folder, filename = os.path.split(os.path.abspath(sheet1_path))
    stem, ext = os.path.splitext(filename)
    backup_path = os.path.join(folder, f"{stem}_backup_{datetime.now():%Y%m%d-%H%M%S}{ext}")
    shutil.copy2(sheet1_path, backup_path)
    try:
        wb.save(sheet1_path)
    except PermissionError:
        raise MatcherError(f"Couldn't save:\n{sheet1_path}\n\n"
                           "If it's open in Excel, close it and try again.")

    # --- Step 8: Build the summary shown to the user ----------------------
    lines = ["Sheet 1 has been updated:", f"  {os.path.abspath(sheet1_path)}",
             "Backup of the original:", f"  {backup_path}", "",
             f"Repair orders checked: {counts['checked']}"]
    for label, path, _ in lookups:
        lines.append(f"  Matched on {label} ({os.path.basename(path)}): {matched[label]}")
    lines.append(f"  Not found (yellow): {counts['not_found']}")
    if counts["no_price"]:
        lines.append(f"  Matched but no price on Sheet 1 (yellow): {counts['no_price']}")
    lines += ["", f"Positive (green): {counts['positive']}",
              f"Negative (red): {counts['negative']}",
              f"Zero (*): {counts['zero']}"]
    if counts["notes"]:
        lines += ["", f"{counts['notes']} cell(s) have a note - hover over the red "
                      "triangle in the corner to read it."]
    return "\n".join(lines)


# ===========================================================================
# Starting exe
# ===========================================================================
def run_with_windows():
    """Ask for the files with pop-up file pickers and show the result in a message box."""
    import tkinter as tk
    from tkinter import filedialog, messagebox

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)   # keep dialogs in front
    title = "Repair Order Matcher"
    lookup_types = [("Excel files", "*.xls *.xlsx"), ("All files", "*.*")]

    # Pick the files (Sheet 3 is optional)
    sheet1 = filedialog.askopenfilename(
        parent=root, title="Step 1 of 3: Select SHEET 1 (the .xlsx file to update)",
        filetypes=[("Excel workbook", "*.xlsx *.xlsm"), ("All files", "*.*")])
    if not sheet1:
        return
    sheet2 = filedialog.askopenfilename(parent=root, title="Step 2 of 3: Select SHEET 2",
                                        filetypes=lookup_types)
    if not sheet2:
        return
    sheet3 = filedialog.askopenfilename(parent=root,
                                        title="Step 3 of 3: Select SHEET 3 (Cancel if there isn't one)",
                                        filetypes=lookup_types)
    lookups = [("Sheet 2", sheet2)] + ([("Sheet 3", sheet3)] if sheet3 else [])

    # Run the matcher and report success or failure
    try:
        summary = process(sheet1, lookups,
                          confirm=lambda msg: messagebox.askyesno(title, msg, parent=root))
    except MatcherError as e:
        messagebox.showerror(title, str(e), parent=root)
    except Exception as e:
        messagebox.showerror(title, f"Something went wrong. No changes were saved to Sheet 1.\n\n"
                                    f"{type(e).__name__}: {e}", parent=root)
    else:
        messagebox.showinfo(title, summary, parent=root)
    root.destroy()


def run_from_command_line(args):
    """Run with file paths given as arguments and print the result. Returns an exit code."""
    if len(args) not in (2, 3):
        print(__doc__)
        return 2
    sheet1, *others = args
    lookups = [(f"Sheet {n}", path) for n, path in enumerate(others, start=2)]

    #  print warning if wrong file type selected
    def confirm(message):
        print("WARNING:", message.replace("\n\nIs this the right Sheet 1 file? Continue anyway?", ""))
        return True

    try:
        print(process(sheet1, lookups, confirm))
    except MatcherError as e:
        print("ERROR:", e)
        return 1
    return 0


if __name__ == "__main__":
    # File paths given -> command line mode; nothing given -> pop-up windows
    if len(sys.argv) > 1:
        sys.exit(run_from_command_line(sys.argv[1:]))
    run_with_windows()
